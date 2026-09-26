import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from conexion import conectar

from openpyxl import Workbook

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib.units import mm

from PIL import Image, ImageTk

import re
import os


def crear_clientes(pestaña):

    # ==========================================================
    # VARIABLES PARA IMAGEN
    # ==========================================================

    imagen_original = None
    imagen_tk = None
    ruta_imagen = ""

    # ==========================================================
    # TÍTULO
    # ==========================================================

    ttk.Label(
        pestaña,
        text="Gestión de Clientes",
        font=("Arial", 16, "bold")
    ).grid(
        row=0,
        column=0,
        columnspan=3,
        pady=15
    )

    # ==========================================================
    # CAMPOS
    # ==========================================================

    campos = [
        ("Código:", "codigo"),
        ("Razón social / Nombre:", "razon_social"),
        ("RUC / Documento:", "ruc"),
        ("Dirección fiscal:", "direccion_fiscal"),
        ("Teléfono:", "telefono"),
        ("Correo electrónico:", "correo"),
        ("Persona de contacto:", "persona_contacto"),
        ("Límite de crédito:", "limite_credito"),
        ("Condiciones de pago:", "condiciones_pago"),
        ("Categoría:", "categoria")
    ]

    entradas = {}

    for fila, (texto, nombre) in enumerate(
        campos,
        start=1
    ):

        ttk.Label(
            pestaña,
            text=texto
        ).grid(
            row=fila,
            column=0,
            padx=10,
            pady=5,
            sticky="e"
        )

        entrada = ttk.Entry(
            pestaña,
            width=35
        )

        entrada.grid(
            row=fila,
            column=1,
            padx=10,
            pady=5
        )

        entradas[nombre] = entrada

    # ==========================================================
    # ÁREA DE IMAGEN
    # ==========================================================

    ttk.Label(
        pestaña,
        text="Imagen del cliente",
        font=("Arial", 11, "bold")
    ).grid(
        row=1,
        column=2,
        padx=20,
        pady=5
    )

    marco_imagen = ttk.Frame(
        pestaña,
        width=180,
        height=150
    )

    marco_imagen.grid(
        row=2,
        column=2,
        rowspan=6,
        padx=20,
        pady=5
    )

    marco_imagen.grid_propagate(False)

    etiqueta_imagen = ttk.Label(
        marco_imagen,
        text="Sin imagen",
        anchor="center"
    )

    etiqueta_imagen.pack(
        fill="both",
        expand=True
    )

    # ==========================================================
    # SELECCIONAR IMAGEN
    # ==========================================================

    def seleccionar_imagen():

        nonlocal imagen_original
        nonlocal imagen_tk
        nonlocal ruta_imagen

        archivo = filedialog.askopenfilename(
            title="Seleccionar imagen del cliente",
            filetypes=[
                (
                    "Imágenes",
                    "*.jpg *.jpeg *.png *.gif"
                )
            ]
        )

        if not archivo:
            return

        # ------------------------------------------------------
        # VALIDAR EXTENSIÓN
        # ------------------------------------------------------

        extension = os.path.splitext(
            archivo
        )[1].lower()

        extensiones_permitidas = [
            ".jpg",
            ".jpeg",
            ".png",
            ".gif"
        ]

        if extension not in extensiones_permitidas:

            messagebox.showwarning(
                "Imagen",
                "Solo se permiten imágenes JPG, PNG o GIF."
            )

            return

        # ------------------------------------------------------
        # VALIDAR TAMAÑO
        # ------------------------------------------------------

        tamaño = os.path.getsize(
            archivo
        )

        tamaño_maximo = 2 * 1024 * 1024

        if tamaño > tamaño_maximo:

            messagebox.showwarning(
                "Imagen",
                "La imagen no puede superar los 2 MB."
            )

            return

        # ------------------------------------------------------
        # ABRIR IMAGEN CON PILLOW
        # ------------------------------------------------------

        try:

            imagen_original = Image.open(
                archivo
            )

            imagen_original.thumbnail(
                (160, 130)
            )

            imagen_tk = ImageTk.PhotoImage(
                imagen_original
            )

            etiqueta_imagen.config(
                image=imagen_tk,
                text=""
            )

            ruta_imagen = archivo

            messagebox.showinfo(
                "Imagen",
                "Imagen seleccionada correctamente."
            )

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"No se pudo abrir la imagen:\n{e}"
            )

    ttk.Button(
        pestaña,
        text="🖼️ Seleccionar imagen",
        command=seleccionar_imagen
    ).grid(
        row=8,
        column=2,
        padx=20,
        pady=5
    )

    # ==========================================================
    # VALIDACIONES
    # ==========================================================

    def validar_campos():

        codigo = entradas["codigo"].get().strip()
        razon_social = entradas["razon_social"].get().strip()
        ruc = entradas["ruc"].get().strip()
        direccion = entradas["direccion_fiscal"].get().strip()
        telefono = entradas["telefono"].get().strip()
        correo = entradas["correo"].get().strip()
        contacto = entradas["persona_contacto"].get().strip()
        limite = entradas["limite_credito"].get().strip()
        condiciones = entradas["condiciones_pago"].get().strip()
        categoria = entradas["categoria"].get().strip()

        # ------------------------------------------------------
        # CAMPOS OBLIGATORIOS
        # ------------------------------------------------------

        if not all([
            codigo,
            razon_social,
            ruc,
            direccion,
            telefono,
            correo,
            contacto,
            limite,
            condiciones,
            categoria
        ]):

            messagebox.showwarning(
                "Validación",
                "Todos los campos son obligatorios."
            )

            return False

        # ------------------------------------------------------
        # CÓDIGO
        # ------------------------------------------------------

        if not codigo.isdigit():

            messagebox.showwarning(
                "Validación",
                "El código debe contener solamente números."
            )

            return False

        if len(codigo) > 10:

            messagebox.showwarning(
                "Validación",
                "El código no puede superar 10 dígitos."
            )

            return False

        # ------------------------------------------------------
        # RAZÓN SOCIAL
        # ------------------------------------------------------

        if len(razon_social) < 2 or len(razon_social) > 100:

            messagebox.showwarning(
                "Validación",
                "La razón social debe tener entre 2 y 100 caracteres."
            )

            return False

        # ------------------------------------------------------
        # RUC
        # ------------------------------------------------------

        if not ruc.isdigit() or len(ruc) != 11:

            messagebox.showwarning(
                "Validación",
                "El RUC debe contener exactamente 11 números."
            )

            return False

        # ------------------------------------------------------
        # DIRECCIÓN
        # ------------------------------------------------------

        if len(direccion) < 5 or len(direccion) > 150:

            messagebox.showwarning(
                "Validación",
                "La dirección debe tener entre 5 y 150 caracteres."
            )

            return False

        # ------------------------------------------------------
        # TELÉFONO
        # ------------------------------------------------------

        if not telefono.isdigit():

            messagebox.showwarning(
                "Validación",
                "El teléfono debe contener solamente números."
            )

            return False

        if len(telefono) < 7 or len(telefono) > 15:

            messagebox.showwarning(
                "Validación",
                "El teléfono debe tener entre 7 y 15 dígitos."
            )

            return False

        # ------------------------------------------------------
        # CORREO
        # ------------------------------------------------------

        patron_correo = r"^[\w\.-]+@[\w\.-]+\.\w+$"

        if not re.match(
            patron_correo,
            correo
        ):

            messagebox.showwarning(
                "Validación",
                "Ingrese un correo electrónico válido."
            )

            return False

        if len(correo) > 100:

            messagebox.showwarning(
                "Validación",
                "El correo no puede superar 100 caracteres."
            )

            return False

        # ------------------------------------------------------
        # CONTACTO
        # ------------------------------------------------------

        if len(contacto) < 2 or len(contacto) > 100:

            messagebox.showwarning(
                "Validación",
                "La persona de contacto debe tener entre 2 y 100 caracteres."
            )

            return False

        # ------------------------------------------------------
        # LÍMITE DE CRÉDITO
        # ------------------------------------------------------

        try:

            limite_numero = float(
                limite
            )

            if limite_numero < 0:
                raise ValueError

        except ValueError:

            messagebox.showwarning(
                "Validación",
                "El límite de crédito debe ser un número mayor o igual a 0."
            )

            return False

        # ------------------------------------------------------
        # CONDICIONES DE PAGO
        # ------------------------------------------------------

        if len(condiciones) < 2 or len(condiciones) > 100:

            messagebox.showwarning(
                "Validación",
                "Las condiciones de pago deben tener entre 2 y 100 caracteres."
            )

            return False

        # ------------------------------------------------------
        # CATEGORÍA
        # ------------------------------------------------------

        if len(categoria) < 2 or len(categoria) > 50:

            messagebox.showwarning(
                "Validación",
                "La categoría debe tener entre 2 y 50 caracteres."
            )

            return False

        return True

    # ==========================================================
    # OBTENER DATOS
    # ==========================================================

    def obtener_clientes():

        datos = []

        try:

            conexion = conectar()

            if conexion:

                cursor = conexion.cursor()

                cursor.callproc(
                    "sp_clientes_listar"
                )

                for resultado in cursor.stored_results():

                    datos.extend(
                        resultado.fetchall()
                    )

                cursor.close()
                conexion.close()

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"No se pudieron consultar los clientes:\n{e}"
            )

        return datos

    # ==========================================================
    # CARGAR TABLA
    # ==========================================================

    columnas = [
        "Código",
        "Razón social",
        "RUC",
        "Dirección",
        "Teléfono",
        "Correo",
        "Contacto",
        "Límite crédito",
        "Condiciones pago",
        "Categoría"
    ]

    tabla = ttk.Treeview(
        pestaña,
        columns=columnas,
        show="headings",
        height=8
    )

    for columna in columnas:

        tabla.heading(
            columna,
            text=columna
        )

        tabla.column(
            columna,
            width=120
        )

    tabla.grid(
        row=14,
        column=0,
        columnspan=3,
        padx=10,
        pady=15,
        sticky="nsew"
    )

    # ==========================================================
    # MOSTRAR CLIENTES
    # ==========================================================

    def mostrar_clientes(datos=None):

        for item in tabla.get_children():

            tabla.delete(item)

        if datos is None:
            datos = obtener_clientes()

        for cliente in datos:

            tabla.insert(
                "",
                tk.END,
                values=cliente
            )

    # ==========================================================
    # FILTRO POR CATEGORÍA
    # ==========================================================

    ttk.Label(
        pestaña,
        text="Filtrar por categoría:"
    ).grid(
        row=10,
        column=2,
        padx=20,
        pady=5
    )

    combo_filtro = ttk.Combobox(
        pestaña,
        width=25,
        state="readonly"
    )

    combo_filtro.grid(
        row=11,
        column=2,
        padx=20,
        pady=5
    )

    combo_filtro["values"] = [
        "Todas",
        "A",
        "B",
        "C",
        "D"
    ]

    combo_filtro.set(
        "Todas"
    )

    # ==========================================================
    # FILTRAR
    # ==========================================================

    def filtrar_clientes():

        datos = obtener_clientes()

        categoria = combo_filtro.get()

        if categoria == "Todas":

            mostrar_clientes(
                datos
            )

            return

        datos_filtrados = []

        for cliente in datos:

            if len(cliente) > 9:

                if str(
                    cliente[9]
                ) == categoria:

                    datos_filtrados.append(
                        cliente
                    )

        mostrar_clientes(
            datos_filtrados
        )

    # ==========================================================
    # CRUD - INSERTAR
    # ==========================================================

    def insertar_cliente():

        if not validar_campos():
            return

        try:

            conexion = conectar()

            if conexion:

                cursor = conexion.cursor()

                cursor.callproc(
                    "sp_clientes_insertar",
                    (
                        int(
                            entradas["codigo"].get()
                        ),
                        entradas["razon_social"].get(),
                        entradas["ruc"].get(),
                        entradas["direccion_fiscal"].get(),
                        entradas["telefono"].get(),
                        entradas["correo"].get(),
                        entradas["persona_contacto"].get(),
                        float(
                            entradas["limite_credito"].get()
                        ),
                        entradas["condiciones_pago"].get(),
                        entradas["categoria"].get()
                    )
                )

                conexion.commit()

                cursor.close()
                conexion.close()

                messagebox.showinfo(
                    "Clientes",
                    "Cliente registrado correctamente."
                )

                mostrar_clientes()

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"No se pudo registrar el cliente:\n{e}"
            )

    # ==========================================================
    # CRUD - ACTUALIZAR
    # ==========================================================

    def actualizar_cliente():

        if not validar_campos():
            return

        try:

            conexion = conectar()

            if conexion:

                cursor = conexion.cursor()

                cursor.callproc(
                    "sp_clientes_actualizar",
                    (
                        int(
                            entradas["codigo"].get()
                        ),
                        entradas["razon_social"].get(),
                        entradas["ruc"].get(),
                        entradas["direccion_fiscal"].get(),
                        entradas["telefono"].get(),
                        entradas["correo"].get(),
                        entradas["persona_contacto"].get(),
                        float(
                            entradas["limite_credito"].get()
                        ),
                        entradas["condiciones_pago"].get(),
                        entradas["categoria"].get()
                    )
                )

                conexion.commit()

                cursor.close()
                conexion.close()

                messagebox.showinfo(
                    "Clientes",
                    "Cliente actualizado correctamente."
                )

                mostrar_clientes()

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"No se pudo actualizar el cliente:\n{e}"
            )

    # ==========================================================
    # CRUD - ELIMINAR
    # ==========================================================

    def eliminar_cliente():

        codigo = entradas["codigo"].get().strip()

        if not codigo:

            messagebox.showwarning(
                "Validación",
                "Ingrese el código del cliente."
            )

            return

        if not codigo.isdigit():

            messagebox.showwarning(
                "Validación",
                "El código debe contener solamente números."
            )

            return

        confirmar = messagebox.askyesno(
            "Confirmar",
            "¿Desea eliminar este cliente?"
        )

        if not confirmar:
            return

        try:

            conexion = conectar()

            if conexion:

                cursor = conexion.cursor()

                cursor.callproc(
                    "sp_clientes_eliminar",
                    (
                        int(codigo),
                    )
                )

                conexion.commit()

                cursor.close()
                conexion.close()

                messagebox.showinfo(
                    "Clientes",
                    "Cliente eliminado correctamente."
                )

                mostrar_clientes()

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"No se pudo eliminar el cliente:\n{e}"
            )

    # ==========================================================
    # LIMPIAR FORMULARIO
    # ==========================================================

    def limpiar():

        nonlocal imagen_original
        nonlocal imagen_tk
        nonlocal ruta_imagen

        for entrada in entradas.values():

            entrada.delete(
                0,
                tk.END
            )

        imagen_original = None
        imagen_tk = None
        ruta_imagen = ""

        etiqueta_imagen.config(
            image="",
            text="Sin imagen"
        )

    # ==========================================================
    # EXPORTAR EXCEL
    # ==========================================================

    def exportar_excel():

        datos = obtener_clientes()

        categoria = combo_filtro.get()

        if categoria != "Todas":

            datos = [
                cliente
                for cliente in datos
                if len(cliente) > 9
                and str(cliente[9]) == categoria
            ]

        if not datos:

            messagebox.showwarning(
                "Excel",
                "No hay datos para exportar."
            )

            return

        try:

            archivo = filedialog.asksaveasfilename(
                title="Guardar clientes en Excel",
                defaultextension=".xlsx",
                filetypes=[
                    ("Archivo Excel", "*.xlsx")
                ]
            )

            if not archivo:
                return

            libro = Workbook()

            hoja = libro.active

            hoja.title = "Clientes"

            for columna, nombre in enumerate(
                columnas,
                start=1
            ):

                hoja.cell(
                    row=1,
                    column=columna,
                    value=nombre
                )

            for fila, cliente in enumerate(
                datos,
                start=2
            ):

                for columna, valor in enumerate(
                    cliente,
                    start=1
                ):

                    hoja.cell(
                        row=fila,
                        column=columna,
                        value=valor
                    )

            for columna in hoja.columns:

                letra = columna[0].column_letter

                hoja.column_dimensions[
                    letra
                ].width = 22

            libro.save(
                archivo
            )

            messagebox.showinfo(
                "Excel",
                "Los clientes fueron exportados correctamente."
            )

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"No se pudo exportar a Excel:\n{e}"
            )

    # ==========================================================
    # EXPORTAR PDF
    # ==========================================================

    def exportar_pdf():

        datos = obtener_clientes()

        categoria = combo_filtro.get()

        if categoria != "Todas":

            datos = [
                cliente
                for cliente in datos
                if len(cliente) > 9
                and str(cliente[9]) == categoria
            ]

        if not datos:

            messagebox.showwarning(
                "PDF",
                "No hay datos para exportar."
            )

            return

        try:

            archivo = filedialog.asksaveasfilename(
                title="Guardar clientes en PDF",
                defaultextension=".pdf",
                filetypes=[
                    ("Archivo PDF", "*.pdf")
                ]
            )

            if not archivo:
                return

            documento = SimpleDocTemplate(
                archivo,
                pagesize=landscape(A4),
                rightMargin=8 * mm,
                leftMargin=8 * mm,
                topMargin=10 * mm,
                bottomMargin=10 * mm
            )

            datos_pdf = [
                columnas
            ]

            datos_pdf.extend(
                datos
            )

            tabla_pdf = Table(
                datos_pdf,
                repeatRows=1
            )

            tabla_pdf.setStyle(
                TableStyle([
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.grey
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold"
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.black
                    )
                ])
            )

            documento.build(
                [tabla_pdf]
            )

            messagebox.showinfo(
                "PDF",
                "Los clientes fueron exportados correctamente."
            )

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"No se pudo exportar a PDF:\n{e}"
            )

    # ==========================================================
    # BOTONES
    # ==========================================================

    ttk.Button(
        pestaña,
        text="Consultar",
        command=mostrar_clientes
    ).grid(
        row=12,
        column=0,
        padx=5,
        pady=10
    )

    ttk.Button(
        pestaña,
        text="Insertar",
        command=insertar_cliente
    ).grid(
        row=12,
        column=1,
        padx=5,
        pady=10
    )

    ttk.Button(
        pestaña,
        text="Actualizar",
        command=actualizar_cliente
    ).grid(
        row=13,
        column=0,
        padx=5,
        pady=10
    )

    ttk.Button(
        pestaña,
        text="Eliminar",
        command=eliminar_cliente
    ).grid(
        row=13,
        column=1,
        padx=5,
        pady=10
    )

    ttk.Button(
        pestaña,
        text="Limpiar",
        command=limpiar
    ).grid(
        row=12,
        column=2,
        padx=5,
        pady=10
    )

    ttk.Button(
        pestaña,
        text="Filtrar",
        command=filtrar_clientes
    ).grid(
        row=13,
        column=2,
        padx=5,
        pady=10
    )

    ttk.Button(
        pestaña,
        text="📊 Exportar Excel",
        command=exportar_excel
    ).grid(
        row=15,
        column=0,
        padx=5,
        pady=10
    )

    ttk.Button(
        pestaña,
        text="📄 Exportar PDF",
        command=exportar_pdf
    ).grid(
        row=15,
        column=1,
        padx=5,
        pady=10
    )

    # ==========================================================
    # MOSTRAR DATOS AL ABRIR
    # ==========================================================

    mostrar_clientes()
